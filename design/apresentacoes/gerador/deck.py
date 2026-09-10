# -*- coding: utf-8 -*-
"""Kit de slides no design system da Dra. Laice.

Camada fina sobre o python-pptx: paleta, tipografia, capa, divisoria de secao,
slide de conteudo com imagem e rodape. Os dois decks (halux valgo e cinderela)
sao montados em cima disto.
"""
import math
import pathlib

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- paleta
BRONZE       = RGBColor(0xB7, 0x8B, 0x4C)
BRONZE_DARK  = RGBColor(0x9A, 0x72, 0x39)
BRONZE_DEEP  = RGBColor(0x6F, 0x51, 0x28)
CREAM_SOFT   = RGBColor(0xF5, 0xE6, 0xD3)
CREAM        = RGBColor(0xFA, 0xF3, 0xE8)
SURFACE      = RGBColor(0xFF, 0xFF, 0xFF)
SURFACE_SOFT = RGBColor(0xFC, 0xFA, 0xF7)
SURFACE_CREAM= RGBColor(0xF9, 0xF3, 0xEA)
INK          = RGBColor(0x3A, 0x3A, 0x3A)
INK_STRONG   = RGBColor(0x1A, 0x1A, 0x1A)
INK_MUTED    = RGBColor(0x6B, 0x6B, 0x6B)
SAGE         = RGBColor(0x7D, 0x94, 0x89)
SAGE_DARK    = RGBColor(0x5D, 0x74, 0x68)
SAGE_SOFT    = RGBColor(0xE8, 0xF0, 0xEC)
BLUE         = RGBColor(0x4A, 0x7C, 0x8A)
BLUE_SOFT    = RGBColor(0xE0, 0xEC, 0xEF)
WARN         = RGBColor(0xB4, 0x53, 0x09)
WARN_SOFT    = RGBColor(0xFD, 0xF0, 0xE2)

# Poppins/Inter sao webfonts e nao estao instaladas: o PowerPoint substituiria
# na hora da apresentacao. Segoe UI e o par mais proximo disponivel no Windows.
FONT = "Segoe UI"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.62)
CONTENT_TOP = Inches(1.72)
FOOTER_Y = Inches(6.86)

ASSINATURA = "Dra. Laice Gomes Cunha · Ortopedia e Traumatologia — Cirurgia do Pé e Tornozelo"
CREDENCIAL = "CRM/SP 204325 | RQE 100836"


# ---------------------------------------------------------------- utilidades
def _shape_no_line(shape):
    shape.line.fill.background()
    shape.shadow.inherit = False


def rect(slide, x, y, w, h, fill=None, line=None, line_w=Pt(1)):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = line_w
    s.shadow.inherit = False
    return s


def round_rect(slide, x, y, w, h, fill, radius=0.06):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    try:
        s.adjustments[0] = radius
    except Exception:
        pass
    return s


def textbox(slide, x, y, w, h, wrap=True, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size=14, bold=False, color=INK, space_after=6,
         align=PP_ALIGN.LEFT, first=False, line=1.25, italic=False, font=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return p


SYMBOL_FONT = "Segoe UI Symbol"     # tem os simbolos que a Segoe UI nao tem


def bullets(tf, items, size=14, color=INK, marker="•", marker_color=BRONZE,
            space_after=9, first=True, line=1.22):
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if (first and i == 0) else tf.add_paragraph()
        p.space_after = Pt(space_after)
        p.line_spacing = line
        rm = p.add_run()
        rm.text = f"{marker}  "
        rm.font.size = Pt(size)
        rm.font.bold = True
        rm.font.color.rgb = marker_color
        rm.font.name = FONT if marker.isascii() else SYMBOL_FONT
        r = p.add_run()
        r.text = it
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = FONT


# ---------------------------------------------------------------- imagens
_conv_cache = {}


MAX_PX = 1500          # nenhuma imagem entra no arquivo maior que isto


def as_png(src, cache_dir):
    """Prepara a imagem para o pptx: python-pptx nao le webp, e as PNG originais
    da cartilha tem ~2 MB cada. Reduz para MAX_PX e grava JPEG quando nao ha
    transparencia, PNG quando ha."""
    src = pathlib.Path(src)
    key = str(src)
    if key in _conv_cache:
        return _conv_cache[key]
    cache_dir = pathlib.Path(cache_dir)
    cache_dir.mkdir(parents=True, exist_ok=True)

    im = Image.open(src)
    if getattr(im, "is_animated", False):
        im.seek(0)
    if max(im.size) > MAX_PX:
        r = MAX_PX / max(im.size)
        im = im.resize((round(im.width * r), round(im.height * r)), Image.LANCZOS)

    alpha = im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info)
    stem = src.name.replace(".", "_")
    if alpha:
        out = cache_dir / f"{stem}.png"
        if not out.exists():
            im.convert("RGBA").save(out, "PNG", optimize=True)
    else:
        out = cache_dir / f"{stem}.jpg"
        if not out.exists():
            im.convert("RGB").save(out, "JPEG", quality=86, optimize=True, progressive=True)
    _conv_cache[key] = out
    return out


def medir_fit(path, w, h, cache_dir):
    """Tamanho que a imagem tera dentro da caixa, no encaixe 'contain'."""
    p = as_png(path, cache_dir)
    iw, ih = Image.open(p).size
    if iw / ih >= w / h:
        return Emu(int(w)), Emu(int(w * ih / iw))
    return Emu(int(h * iw / ih)), Emu(int(h))


def picture_fit(slide, path, x, y, w, h, cache_dir):
    """Encaixa a imagem inteira dentro da caixa, centralizada (contain)."""
    p = as_png(path, cache_dir)
    iw, ih = Image.open(p).size
    box_r, img_r = w / h, iw / ih
    if img_r >= box_r:
        nw, nh = w, Emu(int(w / img_r))
    else:
        nh, nw = h, Emu(int(h * img_r))
    return slide.shapes.add_picture(str(p), Emu(int(x + (w - nw) / 2)),
                                    Emu(int(y + (h - nh) / 2)), nw, nh)


# Largura media de caractere na Segoe UI, em fracao do corpo. Usado so para
# estimar quantas linhas um texto ocupa e escolher o corpo que cabe na caixa.
CHAR_W = 0.53


def cabe(itens, size, largura_in, altura_in, space_after_pt=9, line=1.22,
         topo_in=0.0):
    """Altura estimada, em polegadas, de uma lista de bullets nesse corpo."""
    chars = max(8, int(largura_in / (CHAR_W * size / 72)))
    alt = topo_in
    for it in itens:
        linhas = max(1, -(-len(it) // chars))
        alt += linhas * size * line / 72 + space_after_pt / 72
    return alt, alt <= altura_in


def melhor_corpo(itens, largura_in, altura_in, opcoes, space_after_pt=9,
                 line=1.22, topo_in=0.0):
    """Maior corpo da lista que ainda cabe; o menor se nenhum couber."""
    for size in opcoes:
        if cabe(itens, size, largura_in, altura_in, space_after_pt, line, topo_in)[1]:
            return size
    return opcoes[-1]


# ---------------------------------------------------------------- deck
class Deck:
    def __init__(self, cache_dir, logo=None, assinatura=ASSINATURA):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.cache = pathlib.Path(cache_dir)
        self.logo = logo
        self.assinatura = assinatura
        self.secao = ""

    # -- base
    def _slide(self, bg=SURFACE_SOFT):
        s = self.prs.slides.add_slide(self.blank)
        rect(s, 0, 0, W, H, fill=bg)
        return s

    def _footer(self, s, numero=True):
        rect(s, M, FOOTER_Y, W - 2 * M, Pt(0.9), fill=CREAM_SOFT)
        tf = textbox(s, M, FOOTER_Y + Inches(0.10), Inches(9.2), Inches(0.3))
        para(tf, self.assinatura, size=9, color=INK_MUTED, first=True, space_after=0)
        if numero:
            tf2 = textbox(s, W - M - Inches(2.4), FOOTER_Y + Inches(0.10), Inches(2.4), Inches(0.3))
            para(tf2, str(len(self.prs.slides.__iter__.__self__._sldIdLst) if False else
                          len(self.prs.slides._sldIdLst)),
                 size=9, color=BRONZE, align=PP_ALIGN.RIGHT, first=True, space_after=0)

    def _header(self, s, titulo, kicker=None, sub=None):
        """Desenha o cabecalho e devolve a altura onde o conteudo pode comecar.
        O subtitulo cresce em linhas, entao o resto do slide desce junto."""
        y = Inches(0.62)
        if kicker:
            tf = textbox(s, M, y, W - 2 * M, Inches(0.26))
            para(tf, kicker.upper(), size=10.5, bold=True, color=BRONZE, first=True, space_after=0)
            y += Inches(0.32)
        tw = W - 2 * M - Inches(1.2)
        tf = textbox(s, M, y, tw, Inches(0.62))
        para(tf, titulo, size=27, bold=True, color=INK_STRONG, first=True, space_after=0, line=1.05)
        y += Inches(0.54)
        if sub:
            linhas = max(1, math.ceil(len(sub) / 118))
            tf = textbox(s, M, y + Inches(0.05), tw, Inches(0.24) * linhas)
            para(tf, sub, size=13, color=INK_MUTED, first=True, space_after=0, line=1.2)
            y += Inches(0.05) + Inches(0.23) * linhas
        rect(s, M, y + Inches(0.12), Inches(1.15), Pt(2.6), fill=BRONZE)
        return Emu(int(y + Inches(0.34)))

    # -- tipos de slide
    def capa(self, titulo, subtitulo, linha_paciente="Paciente: ______________________"):
        s = self._slide(bg=CREAM)
        rect(s, 0, 0, Inches(0.34), H, fill=BRONZE)
        rect(s, Inches(7.55), 0, W - Inches(7.55), H, fill=CREAM_SOFT)
        if self.logo:
            picture_fit(s, self.logo, Inches(1.05), Inches(0.85),
                        Inches(2.5), Inches(1.05), self.cache)
        corpo = 40 if len(titulo) <= 22 else (34 if len(titulo) <= 32 else 30)
        tf = textbox(s, Inches(1.05), Inches(2.35), Inches(6.25), Inches(2.0))
        para(tf, titulo, size=corpo, bold=True, color=INK_STRONG, first=True,
             space_after=12, line=1.06)
        para(tf, subtitulo, size=15.5, color=BRONZE_DARK, space_after=0, line=1.3)
        rect(s, Inches(1.05), Inches(4.62), Inches(1.5), Pt(3), fill=BRONZE)
        tf = textbox(s, Inches(1.05), Inches(5.05), Inches(6.0), Inches(1.1))
        para(tf, linha_paciente, size=14, color=INK_MUTED, first=True, space_after=6)
        para(tf, "Data da cirurgia: ____ / ____ / ______", size=14, color=INK_MUTED, space_after=0)
        tf = textbox(s, Inches(1.05), Inches(6.55), Inches(8.0), Inches(0.6))
        para(tf, self.assinatura, size=10.5, color=INK_MUTED, first=True, space_after=2)
        para(tf, CREDENCIAL, size=10.5, bold=True, color=BRONZE, space_after=0)
        return s

    def secao_divisoria(self, numero, titulo, subtitulo=None, imagem=None):
        self.secao = titulo
        s = self._slide(bg=SURFACE_SOFT)
        rect(s, 0, 0, Inches(6.6), H, fill=CREAM_SOFT)
        tf = textbox(s, M, Inches(2.55), Inches(5.4), Inches(0.5))
        para(tf, f"PARTE {numero}", size=12, bold=True, color=BRONZE, first=True, space_after=8)
        tf = textbox(s, M, Inches(3.05), Inches(5.4), Inches(1.5))
        para(tf, titulo, size=34, bold=True, color=INK_STRONG, first=True, space_after=10, line=1.06)
        if subtitulo:
            para(tf, subtitulo, size=14, color=INK_MUTED, space_after=0, line=1.3)
        rect(s, M, Inches(2.18), Inches(1.5), Pt(3), fill=BRONZE)
        if imagem:
            picture_fit(s, imagem, Inches(7.1), Inches(1.0), Inches(5.6), Inches(5.5), self.cache)
        return s

    def conteudo(self, titulo, kicker=None, sub=None, itens=None, imagem=None,
                 nota=None, colunas=None, tags=None, marker="•"):
        """Slide padrao: bullets a esquerda, imagem a direita quando houver."""
        s = self._slide()
        y = self._header(s, titulo, kicker or self.secao, sub)
        col_w = Inches(6.35) if imagem else (W - 2 * M)
        fundo = FOOTER_Y - Inches(0.22)
        if nota:
            fundo = Emu(int(fundo - Inches(0.44)))
        if tags:
            linhas = self._tags_linhas(tags, col_w)
            fundo = Emu(int(fundo - Inches(0.40) * linhas - Inches(0.16)))
        h = Emu(int(fundo - y))

        if colunas:
            self._colunas(s, colunas, M, y, col_w, h)
        elif itens:
            sa = 10 if len(itens) <= 8 else 6
            size = melhor_corpo(itens, Emu(col_w).inches - 0.35, Emu(h).inches,
                                [16, 15, 14, 13, 12, 11, 10], sa)
            tf = textbox(s, M, y, col_w, h)
            bullets(tf, itens, size=size, marker=marker, space_after=sa)

        if tags:
            self._tags(s, tags, M, Emu(int(fundo + Inches(0.16))), col_w)

        if imagem:
            picture_fit(s, imagem, Inches(7.25), Emu(int(y - Inches(0.30))),
                        Inches(5.45), Inches(5.15), self.cache)
        if nota:
            tf = textbox(s, M, FOOTER_Y - Inches(0.52), col_w, Inches(0.4))
            para(tf, nota, size=10.5, italic=True, color=INK_MUTED, first=True, space_after=0)
        self._footer(s)
        return s

    def _colunas(self, s, colunas, x, y, w, h):
        gap = Inches(0.32)
        cw = Emu(int((w - gap * (len(colunas) - 1)) / len(colunas)))
        for i, col in enumerate(colunas):
            cx = Emu(int(x + i * (cw + gap)))
            tint = col.get("tint", SURFACE)
            round_rect(s, cx, y, cw, h, tint)
            bar = col.get("accent", BRONZE)
            rect(s, cx, y, cw, Pt(3.2), fill=bar)
            tit = col["titulo"].upper()
            tf = textbox(s, cx + Inches(0.24), y + Inches(0.26), cw - Inches(0.48), Inches(0.34))
            para(tf, tit, size=11.5 if len(tit) < 34 else 10, bold=True, color=bar,
                 first=True, space_after=0, line=1.15)
            itens = col.get("itens", [])
            cw_in = Emu(cw).inches - 0.48 - 0.30       # menos padding e marcador
            ch_in = Emu(h).inches - 0.92
            size = melhor_corpo(itens, cw_in, ch_in, [13.5, 13, 12, 11, 10.5, 10, 9.5, 9], 8)
            tf = textbox(s, cx + Inches(0.24), y + Inches(0.68), cw - Inches(0.48), h - Inches(0.9))
            bullets(tf, itens, size=size, marker=col.get("marker", "•"),
                    marker_color=bar, space_after=8)

    @staticmethod
    def _tags_linhas(tags, w):
        linhas, cur = 1, 0
        for t in tags:
            tw = Inches(0.22 + 0.085 * len(t))
            if cur + tw > w:
                linhas += 1
                cur = 0
            cur += tw + Inches(0.12)
        return linhas

    def _tags(self, s, tags, x, y, w):
        cur_x, cur_y = x, y
        for t in tags:
            tw = Inches(0.22 + 0.085 * len(t))
            if cur_x + tw > x + w:
                cur_x, cur_y = x, cur_y + Inches(0.42)
            round_rect(s, cur_x, cur_y, tw, Inches(0.34), SAGE_SOFT, radius=0.5)
            tf = textbox(s, cur_x, cur_y + Inches(0.06), tw, Inches(0.24))
            para(tf, t, size=10.5, color=SAGE_DARK, align=PP_ALIGN.CENTER, first=True, space_after=0)
            cur_x = Emu(int(cur_x + tw + Inches(0.12)))

    def exercicio(self, n, total, titulo, categoria, descricao, dose, imagem, grupo=None):
        s = self._slide()
        topo = self._header(s, titulo, grupo or "Exercícios", categoria)
        tf = textbox(s, W - M - Inches(1.6), Inches(0.66), Inches(1.6), Inches(0.34))
        para(tf, f"{n} / {total}", size=12, bold=True, color=BRONZE,
             align=PP_ALIGN.RIGHT, first=True, space_after=0)
        alt = FOOTER_Y - topo - Inches(0.25)
        if imagem:
            # o cartao acompanha a imagem: as fotos sao retrato e uma caixa fixa
            # deixaria uma faixa branca larga ao lado
            pad = Inches(0.14)
            iw, ih = medir_fit(imagem, Inches(6.55) - 2 * pad, alt - 2 * pad, self.cache)
            cw, ch = Emu(int(iw + 2 * pad)), Emu(int(ih + 2 * pad))
            cx = Emu(int(M + (Inches(6.55) - cw) / 2))
            cy = Emu(int(topo + (alt - ch) / 2))
            round_rect(s, cx, cy, cw, ch, SURFACE)
            picture_fit(s, imagem, cx + pad, cy + pad, iw, ih, self.cache)
        tx = M + Inches(6.95)
        tw = W - M - tx
        tf = textbox(s, tx, topo + Inches(0.15), tw, alt - Inches(1.5))
        para(tf, descricao, size=16, color=INK, first=True, space_after=0, line=1.35)
        dose_y = topo + alt - Inches(1.05)
        round_rect(s, tx, dose_y, tw, Inches(1.05), CREAM_SOFT)
        tf = textbox(s, tx + Inches(0.28), dose_y + Inches(0.17), tw - Inches(0.56), Inches(0.8))
        para(tf, "DOSE", size=10, bold=True, color=BRONZE_DARK, first=True, space_after=4)
        para(tf, dose, size=17, bold=True, color=INK_STRONG, space_after=0)
        self._footer(s)
        return s

    def alerta(self, titulo, itens, nota=None, imagem=None):
        s = self._slide(bg=WARN_SOFT)
        topo = self._header(s, titulo, "Atenção")
        w = Inches(6.35) if imagem else (W - 2 * M)
        fundo = FOOTER_Y - Inches(0.25) - (Inches(0.44) if nota else 0)
        round_rect(s, M, topo, w, Emu(int(fundo - topo)), SURFACE)
        tf = textbox(s, M + Inches(0.4), topo + Inches(0.35), w - Inches(0.8),
                     Emu(int(fundo - topo - Inches(0.6))))
        bullets(tf, itens, size=15 if len(itens) <= 7 else 13,
                marker="!", marker_color=WARN, space_after=12)
        if imagem:
            picture_fit(s, imagem, Inches(7.25), topo, Inches(5.45),
                        Emu(int(fundo - topo)), self.cache)
        if nota:
            tf = textbox(s, M, FOOTER_Y - Inches(0.5), W - 2 * M, Inches(0.4))
            para(tf, nota, size=11, bold=True, color=WARN, first=True, space_after=0)
        self._footer(s)
        return s

    def perguntas(self, titulo, pares, kicker="Dúvidas frequentes"):
        s = self._slide()
        y = self._header(s, titulo, kicker)
        h = (FOOTER_Y - y - Inches(0.2) - Inches(0.24) * (len(pares) - 1)) / len(pares)
        for q, a in pares:
            round_rect(s, M, y, W - 2 * M, Emu(int(h)), SURFACE)
            rect(s, M, y, Pt(3.2), Emu(int(h)), fill=BRONZE)
            tf = textbox(s, M + Inches(0.34), y + Inches(0.20), W - 2 * M - Inches(0.68), Emu(int(h)) - Inches(0.3))
            para(tf, q, size=14.5, bold=True, color=INK_STRONG, first=True, space_after=6, line=1.2)
            para(tf, a, size=12.5, color=INK, space_after=0, line=1.28)
            y = Emu(int(y + h + Inches(0.24)))
        self._footer(s)
        return s

    def encerramento(self, titulo, linhas, contato=None):
        s = self._slide(bg=CREAM)
        rect(s, 0, 0, Inches(0.34), H, fill=BRONZE)
        if self.logo:
            picture_fit(s, self.logo, Inches(1.05), Inches(0.8), Inches(2.3), Inches(0.95), self.cache)
        tf = textbox(s, Inches(1.05), Inches(2.3), Inches(7.6), Inches(0.8))
        para(tf, titulo, size=32, bold=True, color=INK_STRONG, first=True, space_after=0, line=1.06)
        rect(s, Inches(1.05), Inches(3.15), Inches(1.5), Pt(3), fill=BRONZE)
        tf = textbox(s, Inches(1.05), Inches(3.55), Inches(8.4), Inches(2.2))
        bullets(tf, linhas, size=15, space_after=12)
        if contato:
            round_rect(s, Inches(9.0), Inches(2.3), Inches(3.5), Inches(2.6), SURFACE)
            tf = textbox(s, Inches(9.3), Inches(2.6), Inches(2.9), Inches(2.2))
            para(tf, "CONTATO", size=10.5, bold=True, color=BRONZE, first=True, space_after=10)
            for c in contato:
                para(tf, c, size=13, color=INK, space_after=8, line=1.25)
        tf = textbox(s, Inches(1.05), Inches(6.35), Inches(9.0), Inches(0.7))
        para(tf, self.assinatura, size=10.5, color=INK_MUTED, first=True, space_after=2)
        para(tf, CREDENCIAL, size=10.5, bold=True, color=BRONZE, space_after=0)
        return s

    def salvar(self, path):
        path = pathlib.Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path
