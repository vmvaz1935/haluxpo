# -*- coding: utf-8 -*-
"""Renderiza um .pptx para PNG lendo as formas reais do arquivo.

Nao substitui o PowerPoint, mas mostra geometria, cores, imagens e quebra de
linha o suficiente para achar sobreposicao e texto estourando da caixa - que e
o que precisa ser conferido antes de apresentar.

  python preview.py <arquivo.pptx> <pasta_saida> [--sheet]
"""
import pathlib
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

DPI = 100                      # 13.333in -> 1333 px
FONTS = pathlib.Path("C:/Windows/Fonts")
_font_cache = {}


def font(size_pt, bold=False, italic=False):
    key = (round(size_pt, 1), bold, italic)
    if key in _font_cache:
        return _font_cache[key]
    name = "segoeui"
    if bold and italic:
        name += "z"
    elif bold:
        name += "b"
    elif italic:
        name += "i"
    px = max(6, round(size_pt * DPI / 72))
    try:
        f = ImageFont.truetype(str(FONTS / f"{name}.ttf"), px)
    except Exception:
        f = ImageFont.load_default()
    _font_cache[key] = f
    return f


def px(emu):
    return round(Emu(emu).inches * DPI)


def rgb(color, default=(0, 0, 0)):
    try:
        if color and color.type is not None and color.rgb is not None:
            return tuple(color.rgb)
    except Exception:
        pass
    return default


def wrap(draw, text, f, max_w):
    linhas, atual = [], ""
    for palavra in text.split():
        teste = (atual + " " + palavra).strip()
        if draw.textlength(teste, font=f) <= max_w or not atual:
            atual = teste
        else:
            linhas.append(atual)
            atual = palavra
    if atual:
        linhas.append(atual)
    return linhas


def render(prs, idx, slide):
    W, H = px(prs.slide_width), px(prs.slide_height)
    img = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(img)
    estouros = []

    for sh in slide.shapes:
        x, y = px(sh.left or 0), px(sh.top or 0)
        w, h = px(sh.width or 0), px(sh.height or 0)

        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                from io import BytesIO
                im = Image.open(BytesIO(sh.image.blob))
                im = im.resize((max(1, w), max(1, h)), Image.LANCZOS)
                if im.mode in ("RGBA", "LA", "P"):
                    im = im.convert("RGBA")
                    img.paste(im.convert("RGB"), (x, y), im.getchannel("A"))
                else:
                    img.paste(im.convert("RGB"), (x, y))
            except Exception:
                d.rectangle([x, y, x + w, y + h], outline=(200, 200, 200))
            continue

        if sh.has_text_frame and sh.text_frame.text.strip() == "" and sh.shape_type is not None:
            try:
                fill = sh.fill
                if fill.type is not None and fill.type == 1:
                    d.rectangle([x, y, x + w, y + h], fill=rgb(fill.fore_color, (240, 240, 240)))
                    continue
            except Exception:
                pass

        if sh.has_text_frame:
            tf = sh.text_frame
            cy = y
            for p in tf.paragraphs:
                runs = [r for r in p.runs if r.text]
                if not runs:
                    cy += 6
                    continue
                size = runs[0].font.size.pt if runs[0].font.size else 14
                bold = bool(runs[0].font.bold)
                ital = bool(runs[0].font.italic)
                col = rgb(runs[0].font.color, (58, 58, 58))
                f = font(size, bold, ital)
                texto = "".join(r.text for r in runs)
                lh = round(size * DPI / 72 * (p.line_spacing or 1.2))
                align = str(p.alignment or "")
                for ln in wrap(d, texto, f, max(10, w)):
                    tx = x
                    if "CENTER" in align:
                        tx = x + (w - d.textlength(ln, font=f)) / 2
                    elif "RIGHT" in align:
                        tx = x + w - d.textlength(ln, font=f)
                    d.text((tx, cy), ln, font=f, fill=col)
                    cy += lh
                cy += round((p.space_after.pt if p.space_after else 6) * DPI / 72)
            if cy > y + h + 4:
                estouros.append((round((cy - (y + h)) / DPI * 72), tf.text[:48]))

    return img, estouros


def main():
    src = pathlib.Path(sys.argv[1])
    out = pathlib.Path(sys.argv[2])
    out.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(src))
    imgs, problemas = [], []
    for i, s in enumerate(prs.slides, 1):
        im, est = render(prs, i, s)
        im.save(out / f"slide-{i:02d}.png")
        imgs.append(im)
        for pts, txt in est:
            problemas.append((i, pts, txt))

    print(f"{len(imgs)} slides renderizados em {out}")
    if problemas:
        print(f"\nCaixas com texto passando do limite ({len(problemas)}):")
        for i, pts, txt in sorted(problemas, key=lambda x: -x[1])[:25]:
            print(f"  slide {i:>2}  +{pts:>3}pt  {txt}")
    else:
        print("nenhuma caixa estourando")

    if "--sheet" in sys.argv:
        cols, cw = 4, 420
        ch = round(cw * imgs[0].height / imgs[0].width)
        rows = (len(imgs) + cols - 1) // cols
        sheet = Image.new("RGB", (cols * (cw + 8) + 8, rows * (ch + 26) + 8), (235, 233, 228))
        dd = ImageDraw.Draw(sheet)
        for i, im in enumerate(imgs):
            r, c = divmod(i, cols)
            X, Y = 8 + c * (cw + 8), 8 + r * (ch + 26)
            sheet.paste(im.resize((cw, ch), Image.LANCZOS), (X, Y))
            dd.text((X + 2, Y + ch + 4), f"{i + 1}", font=font(9), fill=(60, 60, 60))
        sheet.save(out / "_contato.jpg", quality=88)
        print("folha de contato:", out / "_contato.jpg")


if __name__ == "__main__":
    main()
